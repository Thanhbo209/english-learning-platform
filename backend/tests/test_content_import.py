import io

import openpyxl
import pytest
from docx import Document as DocxDocument

from app.services.content_import.importers import (
    FileImportError,
    RawContent,
    import_csv,
    import_docx,
    import_pdf,
    import_xlsx,
)
from app.services.content_import.normalizers import (
    DocumentData,
    ExerciseQuestionData,
    NormalizationError,
    VocabularyItemData,
    normalize_document,
    normalize_exercise,
    normalize_vocabulary,
)
from app.services.content_import.validators import (
    validate_document,
    validate_exercise,
    validate_vocabulary,
)


def _xlsx_bytes(rows: list[list[object]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx_bytes(paragraphs: list[str]) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- Importers ---


def test_import_csv_produces_rows() -> None:
    raw = import_csv(b"word,definition\nhello,a greeting\n")
    assert raw.rows == [{"word": "hello", "definition": "a greeting"}]


def test_import_xlsx_produces_rows() -> None:
    raw = import_xlsx(_xlsx_bytes([["word", "definition"], ["hello", "a greeting"]]))
    assert raw.rows == [{"word": "hello", "definition": "a greeting"}]


def test_import_docx_produces_text() -> None:
    raw = import_docx(_docx_bytes(["First paragraph.", "Second paragraph."]))
    assert raw.text == "First paragraph.\nSecond paragraph."


def test_import_docx_rejects_malformed_file() -> None:
    with pytest.raises(FileImportError):
        import_docx(b"not a real docx file")


def test_import_xlsx_rejects_malformed_file() -> None:
    with pytest.raises(FileImportError):
        import_xlsx(b"not a real xlsx file")


def test_import_pdf_rejects_malformed_file() -> None:
    with pytest.raises(FileImportError):
        import_pdf(b"not a real pdf file")


# --- Normalizers ---


def test_normalize_vocabulary_from_rows() -> None:
    raw = RawContent(
        rows=[{"word": "hello", "definition": "a greeting", "translation": "xin chao"}]
    )
    items = normalize_vocabulary(raw)
    assert items == [
        VocabularyItemData(word="hello", definition="a greeting", translation="xin chao")
    ]


def test_normalize_vocabulary_from_text_lines() -> None:
    raw = RawContent(text="hello: a greeting\nbye: a farewell")
    items = normalize_vocabulary(raw)
    assert [(item.word, item.definition) for item in items] == [
        ("hello", "a greeting"),
        ("bye", "a farewell"),
    ]


def test_normalize_vocabulary_raises_when_empty() -> None:
    with pytest.raises(NormalizationError):
        normalize_vocabulary(RawContent())


def test_normalize_exercise_from_rows() -> None:
    raw = RawContent(
        rows=[
            {
                "question": "2+2?",
                "type": "multiple_choice",
                "option_a": "3",
                "option_b": "4",
                "answer": "4",
            }
        ]
    )
    questions = normalize_exercise(raw)
    assert questions[0].question_text == "2+2?"
    assert questions[0].options == ["3", "4"]
    assert questions[0].correct_answer == "4"


def test_normalize_exercise_rejects_text_only_input() -> None:
    with pytest.raises(NormalizationError):
        normalize_exercise(RawContent(text="Question 1: what is 2+2? Answer: 4"))


def test_normalize_exercise_handles_spaced_headers_and_numbered_options() -> None:
    # Real-world exports (e.g. quiz tools) use headers like "Question Text" /
    # "Option 1".."Option 5" / "Correct Answer" and vendor-specific type
    # labels instead of our internal "multiple_choice" value.
    raw = RawContent(
        rows=[
            {
                "Question Text": "2+2?",
                "Question Type": "Quiz",
                "Option 1": "3",
                "Option 2": "4",
                "Option 3": "",
                "Correct Answer": "4",
                "Time in seconds": "20",
            }
        ]
    )
    questions = normalize_exercise(raw)
    assert questions[0].question_text == "2+2?"
    assert questions[0].question_type == "multiple_choice"
    assert questions[0].options == ["3", "4"]
    assert questions[0].correct_answer == "4"


def test_normalize_exercise_resolves_positional_correct_answer() -> None:
    # Real-world export: "Correct Answer" holds a 1-based option position
    # ("2") instead of repeating the option text.
    raw = RawContent(
        rows=[
            {
                "Question Text": "He writes letters to his grandparents every week.",
                "Question Type": "Multiple Choice",
                "Option 1": "Letters are written to his grandparents every week.",
                "Option 2": "Letters are written every week to his grandparents.",
                "Option 3": "Letters were written every week to his grandparents.",
                "Option 4": "Letters write every week to his grandparents.",
                "Correct Answer": "2",
            }
        ]
    )
    questions = normalize_exercise(raw)
    assert questions[0].correct_answer == "Letters are written every week to his grandparents."


def test_normalize_exercise_resolves_lettered_correct_answer() -> None:
    raw = RawContent(
        rows=[{"question": "2+2?", "option_a": "3", "option_b": "4", "answer": "B"}]
    )
    questions = normalize_exercise(raw)
    assert questions[0].correct_answer == "4"


def test_normalize_exercise_falls_back_to_short_answer_without_options() -> None:
    raw = RawContent(
        rows=[{"question text": "Name a fruit.", "question type": "Open Ended", "answer": "apple"}]
    )
    questions = normalize_exercise(raw)
    assert questions[0].question_type == "short_answer"
    assert questions[0].options is None


def test_normalize_document_from_text() -> None:
    data = normalize_document(RawContent(text="  Some document body.  "))
    assert data.body == "Some document body."


def test_normalize_document_raises_when_empty() -> None:
    with pytest.raises(NormalizationError):
        normalize_document(RawContent(text="   "))


# --- Validators ---


def test_validate_vocabulary_flags_missing_fields_and_duplicates() -> None:
    items = [
        VocabularyItemData(word="", definition="a greeting"),
        VocabularyItemData(word="hello", definition=""),
        VocabularyItemData(word="hello", definition="a greeting"),
    ]
    issues = validate_vocabulary(items)
    messages = [issue.message for issue in issues]
    assert "Missing word." in messages
    assert "Missing definition." in messages
    assert any("Duplicate word" in message for message in messages)


def test_validate_vocabulary_passes_for_valid_items() -> None:
    items = [VocabularyItemData(word="hello", definition="a greeting")]
    assert validate_vocabulary(items) == []


def test_validate_exercise_flags_missing_answer_and_short_options() -> None:
    questions = [
        ExerciseQuestionData(
            question_text="2+2?",
            question_type="multiple_choice",
            options=["4"],
            correct_answer="",
        )
    ]
    issues = validate_exercise(questions)
    messages = [issue.message for issue in issues]
    assert "Missing correct answer." in messages
    assert any("fewer than 2 options" in message for message in messages)


def test_validate_exercise_flags_answer_not_matching_options() -> None:
    questions = [
        ExerciseQuestionData(
            question_text="2+2?",
            question_type="multiple_choice",
            options=["3", "5"],
            correct_answer="4",
        )
    ]
    issues = validate_exercise(questions)
    assert any("does not match any of the options" in issue.message for issue in issues)


def test_validate_exercise_passes_for_valid_question() -> None:
    questions = [
        ExerciseQuestionData(
            question_text="2+2?",
            question_type="multiple_choice",
            options=["3", "4"],
            correct_answer="4",
        )
    ]
    assert validate_exercise(questions) == []


def test_validate_document_flags_empty_body() -> None:
    issues = validate_document(DocumentData(body=""))
    assert len(issues) == 1


# --- Phase 1 Extended Behavior Tests ---


def test_import_csv_handles_cp1252_encoding() -> None:
    # "café" in CP1252 non-UTF8 bytes
    cp1252_bytes = "word,definition\ncafé,a coffee shop\n".encode("cp1252")
    raw = import_csv(cp1252_bytes)
    assert raw.rows == [{"word": "café", "definition": "a coffee shop"}]


def test_normalize_vocabulary_vietnamese_headers() -> None:
    raw = RawContent(
        rows=[
            {
                "từ vựng": "apple",
                "giải thích": "quả táo",
                "nghĩa tiếng việt": "táo",
                "câu ví dụ": "I eat an apple.",
            }
        ]
    )
    items = normalize_vocabulary(raw)
    assert items[0] == VocabularyItemData(
        word="apple",
        definition="quả táo",
        translation="táo",
        example="I eat an apple.",
    )


def test_normalize_exercise_standalone_letter_options() -> None:
    raw = RawContent(
        rows=[
            {
                "câu hỏi": "What is 1+1?",
                "A": "1",
                "B": "2",
                "C": "3",
                "D": "4",
                "đáp án": "B",
            }
        ]
    )
    questions = normalize_exercise(raw)
    assert questions[0].question_text == "What is 1+1?"
    assert questions[0].question_type == "multiple_choice"
    assert questions[0].options == ["1", "2", "3", "4"]
    assert questions[0].correct_answer == "2"


def test_normalize_exercise_choice_headers_and_whitespace_answer() -> None:
    raw = RawContent(
        rows=[
            {
                "Nội dung câu hỏi": "Capital of France?",
                "Choice 1": "London",
                "Choice 2": "Paris",
                "Đáp án đúng": "  paris  ",
            }
        ]
    )
    questions = normalize_exercise(raw)
    assert questions[0].options == ["London", "Paris"]
    assert questions[0].correct_answer == "Paris"


# --- Phase 2 Tests: Column Mapping & Structure Detection & Severity ---


def test_detect_structure_and_suggest_mappings_vocabulary() -> None:
    from app.services.content_import.normalizers import detect_structure_and_suggest_mappings

    raw = RawContent(
        rows=[{"Term": "apple", "Meaning": "fruit", "Extra": "info"}]
    )
    result = detect_structure_and_suggest_mappings(raw, "vocabulary")
    assert result["headers"] == ["Term", "Meaning", "Extra"]
    assert result["suggested_mapping"]["word"] == "Term"
    assert result["suggested_mapping"]["definition"] == "Meaning"
    assert "Extra" in result["unrecognized_headers"]
    assert result["missing_required_fields"] == []


def test_normalize_vocabulary_custom_column_mapping() -> None:
    raw = RawContent(
        rows=[{"ColA": "book", "ColB": "sách", "ColC": "reading"}]
    )
    mapping = {"word": "ColA", "definition": "ColB", "example": "ColC"}
    items = normalize_vocabulary(raw, column_mapping=mapping)
    assert items[0] == VocabularyItemData(
        word="book", definition="sách", translation=None, example="reading"
    )


def test_normalize_exercise_custom_column_mapping() -> None:
    raw = RawContent(
        rows=[
            {
                "MyQuestion": "What is 2+2?",
                "Opt1": "3",
                "Opt2": "4",
                "MyAnswer": "4",
            }
        ]
    )
    mapping = {
        "question_text": "MyQuestion",
        "options": ["Opt1", "Opt2"],
        "correct_answer": "MyAnswer",
    }
    questions = normalize_exercise(raw, column_mapping=mapping)
    assert questions[0].question_text == "What is 2+2?"
    assert questions[0].options == ["3", "4"]
    assert questions[0].correct_answer == "4"


def test_validate_vocabulary_severity_error_vs_warning() -> None:
    items = [
        VocabularyItemData(word="cat", definition="mèo"),
        VocabularyItemData(word="cat", definition="con mèo"),  # Duplicate word -> warning
        VocabularyItemData(word="", definition="missing word"),  # Missing word -> error
    ]
    issues = validate_vocabulary(items)
    errors = [i for i in issues if i.severity == "error"]
    warnings = [i for i in issues if i.severity == "warning"]

    assert len(errors) == 1
    assert errors[0].row_index == 3
    assert len(warnings) == 1
    assert warnings[0].row_index == 2


